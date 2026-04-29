import os 
import streamlit as st
import time
from openai import OpenAI
from dotenv import load_dotenv
import PyPDF2
import docx
import re
from fpdf import FPDF # PDF generation ke liye

# ---------------- CONFIG ----------------
st.set_page_config(
    page_title="AI Interviewer",
    page_icon="logo.png",
    layout="wide"
)

# ---------------- PREMIUM UI (SAME AS ORIGINAL) ----------------
st.markdown("""
<style>

/* 🔥 Clean coding Image Background (Balanced & Sharp) */
.stApp {
    background:
        linear-gradient(rgba(0,0,0,0.65), rgba(0,0,0,0.75)),
        url("https://images.unsplash.com/photo-1555949963-aa79dcee981c?q=80&w=1920&auto=format&fit=crop");

    background-size: cover;
    background-position: center;
    background-attachment: fixed;

    color: white;
}

/* Optional soft depth (NO blur, only light overlay) */
.stApp::before {
    content: "";
    position: fixed;
    inset: 0;
    background: rgba(0,0,0,0.15);
    pointer-events: none;
}

/* Layout */
.block-container {
    padding-top: 2rem;
}

/* Hero */
.hero {
    margin-top: 20px;
    font-size: 46px;
    font-weight: 800;
    text-shadow: 0 0 20px rgba(250,204,21,0.4);
    background: linear-gradient(90deg, #facc15, #f59e0b);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
}

/* Sub */
.sub {
    font-size: 16px;
    color: #cbd5e1;
}

/* Cards */
.card {
    background: rgba(255,255,255,0.06);
    backdrop-filter: blur(8px);
    padding: 12px;
    border-radius: 10px;
    border: 1px solid rgba(255,255,255,0.1);
    text-align: center;
    font-size: 13px;
}

/* Login box */
.login-box {
    background: rgba(255,255,255,0.06);
    backdrop-filter: blur(10px);
    padding: 14px;
    border-radius: 10px;
    border: 1px solid rgba(255,255,255,0.1);
}

/* Button */
div.stButton > button {
    background-color: #facc15;
    color: black;
    font-weight: bold;
    border-radius: 8px;
}
div.stButton > button:hover {
    background-color: #eab308;
}

/* Timer */
.timer {
    font-size: 20px;
    font-weight: bold;
    color: #facc15;
}

/* Chat box */
.chat-box {
    background: rgba(255,255,255,0.06);
    padding: 15px;
    border-radius: 10px;
    border: 1px solid rgba(255,255,255,0.1);
}

</style>
""", unsafe_allow_html=True)

# ---------------- LOAD API ----------------
load_dotenv()
client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))

# ---------------- HELPERS ----------------
def extract_text(file):
    text = ""
    if file.type == "application/pdf":
        pdf = PyPDF2.PdfReader(file)
        for page in pdf.pages:
            text += page.extract_text() or ""
    elif file.type == "text/plain":
        text = file.read().decode("utf-8", errors="ignore")
    elif "wordprocessingml" in file.type:
        doc = docx.Document(file)
        for para in doc.paragraphs:
            text += para.text + "\n"
    return text

def generate_questions(text):
    with st.spinner("🤖 Generating questions..."):
        res = client.chat.completions.create(
            model="gpt-4o",
            messages=[
                {"role": "system", "content": "Generate exactly 7 clean interview questions."},
                {"role": "user", "content": text}
            ]
        )
    raw = res.choices[0].message.content.split("\n")
    return [q.strip("-•1234567890. ") for q in raw if q.strip()][:7]

def evaluate(q, a):
    with st.spinner("🤖 Evaluating..."):
        res = client.chat.completions.create(
            model="gpt-4o",
            messages=[
                {"role": "system", "content": "Return exactly this format:\nScore: X/5\nFeedback: [1 short sentence]\nImprovement: [1 short tip]"},
                {"role": "user", "content": f"{q}\n{a}"}
            ]
        )
    return res.choices[0].message.content

def chatbot_res(msg):
    res = client.chat.completions.create(
        model="gpt-4o",
        messages=[{"role": "system", "content": "You are a helpful assistant."}, {"role": "user", "content": msg}]
    )
    return res.choices[0].message.content

def create_pdf(questions, evaluations):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(200, 10, "Interview Performance Report", ln=True, align='C')
    pdf.ln(10)
    for i in range(len(questions)):
        pdf.set_font("Arial", 'B', 12)
        pdf.multi_cell(0, 10, f"Q{i+1}: {questions[i]}")
        pdf.set_font("Arial", '', 11)
        pdf.multi_cell(0, 7, f"{evaluations[i]}")
        pdf.ln(5)
    return pdf.output(dest='S').encode('latin-1')

# ---------------- LOGIN ----------------
APP_PASSWORD = "admin1234"
if "logged" not in st.session_state: st.session_state.logged = False
if "mode" not in st.session_state: st.session_state.mode = "Interview Assistant"

if not st.session_state.logged:
    st.markdown('<div class="hero">Your Next Role Starts Here.</div>', unsafe_allow_html=True)
    st.markdown('<p class="sub">AI Powered Interview Practice</p>', unsafe_allow_html=True)
    st.markdown("<br>", unsafe_allow_html=True)

    # All 6 Cards (Original)
    c1, c2, c3 = st.columns(3)
    with c1: st.markdown('<div class="card">📄 Resume Based</div>', unsafe_allow_html=True)
    with c2: st.markdown('<div class="card">⚡ AI Feedback</div>', unsafe_allow_html=True)
    with c3: st.markdown('<div class="card">📊 Report</div>', unsafe_allow_html=True)
    st.markdown("<br>", unsafe_allow_html=True)
    c4, c5, c6 = st.columns(3)
    with c4: st.markdown('<div class="card">🎯 Smart Questions</div>', unsafe_allow_html=True)
    with c5: st.markdown('<div class="card">🤖 Chat With AI</div>', unsafe_allow_html=True)
    with c6: st.markdown('<div class="card">🚀 AI Evaluation</div>', unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)

    left, right = st.columns([2,1])
    
    with left:
        st.markdown("### 🤖 About AI Interviewer")
        st.markdown("""
🚀 Resume-based smart questions
🚀 Real-time AI evaluation
🚀 Detailed feedback & improvement
🚀 30-minute timed interview
🚀 Instant selection decision
""")

    with right:
        st.markdown('<div class="login-box">', unsafe_allow_html=True)
        with st.form("login_form"):
            email = st.text_input("Email")
            password = st.text_input("Password", type="password")
            if st.form_submit_button("Login"):
                if password == APP_PASSWORD:
                    st.session_state.logged = True
                    st.rerun()
                else: st.error("Wrong password")
        st.markdown('</div>', unsafe_allow_html=True)

# ---------------- MAIN ----------------
else:
    with st.sidebar:
        st.markdown("## 🤖 AI Interviewer")
        st.session_state.mode = st.radio("Select Mode", ["Interview Assistant", "Chatbot"])
        if st.button("🔄 Restart"):
            st.session_state.clear()
            st.rerun()

    # ---------------- CHATBOT (FIXED) ----------------
    if st.session_state.mode == "Chatbot":
        st.markdown('<div class="hero">AI Chatbot</div>', unsafe_allow_html=True)
        if "chat" not in st.session_state: st.session_state.chat = []

        for role, msg in st.session_state.chat:
            st.markdown(f"🧑 {msg}" if role == "user" else f"🤖 {msg}")

        # Auto-clear input logic
        def handle_chat():
            u_msg = st.session_state.chat_input
            if u_msg:
                reply = chatbot_res(u_msg)
                st.session_state.chat.append(("user", u_msg))
                st.session_state.chat.append(("ai", reply))
                st.session_state.chat_input = "" # Input khali ho jayega

        st.text_input("Ask something...", key="chat_input", on_change=handle_chat)

    # ---------------- INTERVIEW ----------------
    else:
        if "step" not in st.session_state:
            st.session_state.step, st.session_state.q, st.session_state.a, st.session_state.e, st.session_state.i = "upload", [], [], [], 0

        if st.session_state.step == "upload":
            st.markdown('<div class="hero">AI Interviewer</div>', unsafe_allow_html=True)
            file = st.file_uploader("Upload Resume", type=["pdf","txt","docx"])
            if file and st.button("Start Interview"):
                text = extract_text(file)
                st.session_state.q = generate_questions(text)
                st.session_state.step, st.session_state.start = "interview", time.time()
                st.rerun()

        elif st.session_state.step == "interview":
            rem = 1800 - int(time.time() - st.session_state.start)
            st.markdown(f'<div class="timer">⏱ {rem//60:02}:{rem%60:02}</div>', unsafe_allow_html=True)
            st.progress((st.session_state.i+1)/7)
            
            q = st.session_state.q[st.session_state.i]
            st.subheader(f"Question {st.session_state.i+1}")
            st.write(q)
            ans = st.text_area("Your Answer", key=f"ans_{st.session_state.i}", height=150)

            if st.button("Submit", use_container_width=True):
                if not ans.strip(): st.warning("Write answer")
                else:
                    st.session_state.e.append(evaluate(q, ans))
                    st.session_state.a.append(ans)
                    if st.session_state.i < 6: st.session_state.i += 1
                    else: st.session_state.step = "report"
                    st.rerun()

        elif st.session_state.step == "report":
            st.markdown("## 📊 Final Report")
            score = 0
            for i, e in enumerate(st.session_state.e):
                st.markdown(f'<div class="chat-box"><b>Question {i+1}</b><br>{e}</div>', unsafe_allow_html=True)
                m = re.search(r"Score:\s*(\d)", e)
                if m and int(m.group(1)) >= 4: score += 1
            
            st.write(f"### Score: {score}/7")
            
            # PDF Download Button
            pdf_data = create_pdf(st.session_state.q, st.session_state.e)
            st.download_button("📥 Download Report PDF", pdf_data, "Report.pdf", "application/pdf")

            if score >= 5: st.success("🎉 Selected")
            else: st.error("❌ Rejected")