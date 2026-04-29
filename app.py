import os
import streamlit as st
import time
import PyPDF2
import docx
import re
from fpdf import FPDF
from openai import OpenAI
from dotenv import load_dotenv

# ---------------- CONFIG ----------------
st.set_page_config(
    page_title="AI Interviewer",
    page_icon="🤖",
    layout="wide"
)

# ---------------- STREAMLIT-SAFE CSS ----------------
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Space+Grotesk:wght@400;600;700&family=Inter:wght@400;500&display=swap');

/* ── Global reset ── */
html, body, [class*="css"] {
    font-family: 'Inter', sans-serif;
    color: #e2e8f0;
}

.stApp {
    background: linear-gradient(135deg, #0f172a 0%, #1e1b4b 50%, #0f172a 100%);
    min-height: 100vh;
}

/* ── Fix ALL input text visibility ── */
input[type="text"],
input[type="password"],
input[type="email"],
textarea,
.stTextInput > div > div > input,
.stTextArea > div > div > textarea {
    background-color: #1e293b !important;
    color: #f1f5f9 !important;
    border: 1.5px solid #334155 !important;
    border-radius: 10px !important;
    padding: 10px 14px !important;
    font-size: 15px !important;
    caret-color: #facc15 !important;
}

input[type="text"]:focus,
input[type="password"]:focus,
.stTextInput > div > div > input:focus,
.stTextArea > div > div > textarea:focus {
    border-color: #facc15 !important;
    box-shadow: 0 0 0 3px rgba(250, 204, 21, 0.15) !important;
    outline: none !important;
}

/* ── Label text ── */
.stTextInput > label,
.stTextArea > label,
.stFileUploader > label {
    color: #94a3b8 !important;
    font-size: 13px !important;
    font-weight: 500 !important;
    letter-spacing: 0.05em;
    text-transform: uppercase;
    margin-bottom: 4px !important;
}

/* ── Placeholder text ── */
input::placeholder,
textarea::placeholder {
    color: #475569 !important;
}

/* ── Buttons ── */
div.stButton > button {
    background: linear-gradient(135deg, #facc15, #f59e0b) !important;
    color: #0f172a !important;
    font-weight: 700 !important;
    font-size: 15px !important;
    border: none !important;
    border-radius: 10px !important;
    padding: 10px 28px !important;
    cursor: pointer;
    transition: all 0.2s ease;
    letter-spacing: 0.03em;
}
div.stButton > button:hover {
    background: linear-gradient(135deg, #fde047, #facc15) !important;
    transform: translateY(-1px);
    box-shadow: 0 6px 20px rgba(250, 204, 21, 0.35) !important;
}
div.stButton > button:active {
    transform: translateY(0);
}

/* ── Download button ── */
div.stDownloadButton > button {
    background: linear-gradient(135deg, #3b82f6, #2563eb) !important;
    color: white !important;
    font-weight: 700 !important;
    border: none !important;
    border-radius: 10px !important;
    padding: 10px 28px !important;
}

/* ── Radio buttons ── */
.stRadio > div {
    background: rgba(255,255,255,0.04);
    border-radius: 10px;
    padding: 8px;
}
.stRadio label {
    color: #cbd5e1 !important;
}

/* ── Progress bar ── */
.stProgress > div > div {
    background: linear-gradient(90deg, #facc15, #f59e0b) !important;
    border-radius: 10px !important;
}
.stProgress > div {
    background: #1e293b !important;
    border-radius: 10px !important;
}

/* ── Sidebar ── */
section[data-testid="stSidebar"] {
    background: linear-gradient(180deg, #0f172a, #1e1b4b) !important;
    border-right: 1px solid rgba(255,255,255,0.07);
}
section[data-testid="stSidebar"] * {
    color: #e2e8f0 !important;
}

/* ── Alerts / Messages ── */
.stSuccess {
    background: rgba(34,197,94,0.15) !important;
    border: 1px solid rgba(34,197,94,0.4) !important;
    border-radius: 10px !important;
    color: #86efac !important;
}
.stError {
    background: rgba(239,68,68,0.15) !important;
    border: 1px solid rgba(239,68,68,0.4) !important;
    border-radius: 10px !important;
    color: #fca5a5 !important;
}
.stWarning {
    background: rgba(245,158,11,0.15) !important;
    border: 1px solid rgba(245,158,11,0.4) !important;
    border-radius: 10px !important;
}

/* ── File uploader ── */
[data-testid="stFileUploader"] {
    background: rgba(255,255,255,0.04) !important;
    border: 2px dashed #334155 !important;
    border-radius: 12px !important;
    padding: 20px !important;
}
[data-testid="stFileUploader"]:hover {
    border-color: #facc15 !important;
}

/* ── Spinner ── */
.stSpinner > div {
    border-top-color: #facc15 !important;
}

/* ── Hide streamlit branding ── */
#MainMenu, footer, header { visibility: hidden; }
.block-container { padding-top: 2rem !important; padding-bottom: 2rem !important; }

/* ── Custom components ── */
.hero-title {
    font-family: 'Space Grotesk', sans-serif;
    font-size: 52px;
    font-weight: 700;
    background: linear-gradient(90deg, #facc15 0%, #f97316 100%);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
    background-clip: text;
    line-height: 1.1;
    margin-bottom: 8px;
}
.hero-sub {
    font-size: 17px;
    color: #94a3b8;
    margin-bottom: 32px;
    letter-spacing: 0.02em;
}
.feature-card {
    background: rgba(255,255,255,0.05);
    border: 1px solid rgba(255,255,255,0.09);
    border-radius: 14px;
    padding: 18px 16px;
    text-align: center;
    transition: all 0.2s;
}
.feature-card:hover {
    border-color: rgba(250,204,21,0.4);
    background: rgba(250,204,21,0.05);
}
.feature-icon { font-size: 24px; margin-bottom: 6px; }
.feature-label { font-size: 13px; color: #94a3b8; font-weight: 500; }
.login-card {
    background: rgba(255,255,255,0.04);
    border: 1px solid rgba(255,255,255,0.1);
    border-radius: 16px;
    padding: 28px 24px;
}
.login-title {
    font-family: 'Space Grotesk', sans-serif;
    font-size: 20px;
    font-weight: 600;
    color: #f1f5f9;
    margin-bottom: 4px;
}
.login-sub {
    font-size: 13px;
    color: #64748b;
    margin-bottom: 20px;
}
.section-divider {
    border: none;
    border-top: 1px solid rgba(255,255,255,0.07);
    margin: 24px 0;
}
.timer-badge {
    display: inline-flex;
    align-items: center;
    gap: 8px;
    background: rgba(250,204,21,0.12);
    border: 1px solid rgba(250,204,21,0.3);
    border-radius: 100px;
    padding: 8px 20px;
    font-size: 22px;
    font-weight: 700;
    font-family: 'Space Grotesk', sans-serif;
    color: #facc15;
    margin-bottom: 16px;
}
.question-card {
    background: rgba(255,255,255,0.04);
    border: 1px solid rgba(255,255,255,0.1);
    border-left: 4px solid #facc15;
    border-radius: 12px;
    padding: 20px;
    margin-bottom: 16px;
}
.question-num {
    font-size: 12px;
    font-weight: 600;
    color: #facc15;
    text-transform: uppercase;
    letter-spacing: 0.1em;
    margin-bottom: 6px;
}
.question-text {
    font-size: 17px;
    color: #f1f5f9;
    font-weight: 500;
    line-height: 1.5;
}
.eval-card {
    background: rgba(255,255,255,0.04);
    border: 1px solid rgba(255,255,255,0.09);
    border-radius: 12px;
    padding: 20px;
    margin-bottom: 16px;
}
.eval-q-label {
    font-size: 12px;
    font-weight: 600;
    color: #facc15;
    text-transform: uppercase;
    letter-spacing: 0.08em;
    margin-bottom: 8px;
}
.score-display {
    font-family: 'Space Grotesk', sans-serif;
    font-size: 42px;
    font-weight: 700;
    color: #facc15;
    margin: 8px 0 4px;
}
.chat-bubble-user {
    background: rgba(250,204,21,0.1);
    border: 1px solid rgba(250,204,21,0.2);
    border-radius: 12px 12px 4px 12px;
    padding: 12px 16px;
    margin: 8px 0 8px 20%;
    color: #fef3c7;
    font-size: 15px;
}
.chat-bubble-ai {
    background: rgba(255,255,255,0.05);
    border: 1px solid rgba(255,255,255,0.1);
    border-radius: 12px 12px 12px 4px;
    padding: 12px 16px;
    margin: 8px 20% 8px 0;
    color: #e2e8f0;
    font-size: 15px;
}
.about-item {
    display: flex;
    align-items: center;
    gap: 10px;
    padding: 8px 0;
    color: #cbd5e1;
    font-size: 15px;
    border-bottom: 1px solid rgba(255,255,255,0.05);
}
</style>
""", unsafe_allow_html=True)

# ---------------- LOAD API ----------------
load_dotenv()
client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))

# ---------------- HELPERS ----------------
def extract_text(file):
    text = ""
    if file.type == "application/pdf":
        pdf = PyPDF2.PdfReader(file)
        for page in pdf.pages:
            text += page.extract_text() or ""
    elif file.type == "text/plain":
        text = file.read().decode("utf-8", errors="ignore")
    elif "wordprocessingml" in file.type:
        doc = docx.Document(file)
        for para in doc.paragraphs:
            text += para.text + "\n"
    return text

def generate_questions(text):
    with st.spinner("Generating interview questions..."):
        res = client.chat.completions.create(
            model="gpt-4o",
            messages=[
                {"role": "system", "content": "Generate exactly 7 clean interview questions."},
                {"role": "user", "content": text}
            ]
        )
    raw = res.choices[0].message.content.split("\n")
    return [q.strip("-•1234567890. ") for q in raw if q.strip()][:7]

def evaluate(q, a):
    with st.spinner("Evaluating your answer..."):
        res = client.chat.completions.create(
            model="gpt-4o",
            messages=[
                {"role": "system", "content": "Return exactly this format:\nScore: X/5\nFeedback: [1 short sentence]\nImprovement: [1 short tip]"},
                {"role": "user", "content": f"{q}\n{a}"}
            ]
        )
    return res.choices[0].message.content

def chatbot_res(msg):
    res = client.chat.completions.create(
        model="gpt-4o",
        messages=[
            {"role": "system", "content": "You are a helpful AI assistant for interview preparation."},
            {"role": "user", "content": msg}
        ]
    )
    return res.choices[0].message.content

def create_pdf(questions, evaluations):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(200, 10, "Interview Performance Report", ln=True, align='C')
    pdf.ln(10)
    for i in range(len(questions)):
        pdf.set_font("Arial", 'B', 12)
        pdf.multi_cell(0, 10, f"Q{i+1}: {questions[i]}")
        pdf.set_font("Arial", '', 11)
        pdf.multi_cell(0, 7, f"{evaluations[i]}")
        pdf.ln(5)
    return pdf.output(dest='S').encode('latin-1')

# ---------------- SESSION INIT ----------------
APP_PASSWORD = "admin1234"
if "logged" not in st.session_state: st.session_state.logged = False
if "mode" not in st.session_state: st.session_state.mode = "Interview Assistant"

# ================================================================
# LOGIN PAGE
# ================================================================
if not st.session_state.logged:
    st.markdown('<div class="hero-title">Your Next Role<br>Starts Here.</div>', unsafe_allow_html=True)
    st.markdown('<div class="hero-sub">AI-Powered Mock Interview Platform</div>', unsafe_allow_html=True)

    # Feature cards
    cols = st.columns(6)
    features = [
        ("📄", "Resume Based"),
        ("⚡", "AI Feedback"),
        ("📊", "PDF Report"),
        ("🎯", "Smart Questions"),
        ("🤖", "AI Chatbot"),
        ("🚀", "Live Scoring"),
    ]
    for col, (icon, label) in zip(cols, features):
        with col:
            st.markdown(f"""
            <div class="feature-card">
                <div class="feature-icon">{icon}</div>
                <div class="feature-label">{label}</div>
            </div>""", unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)
    left, _, right = st.columns([1.1, 0.1, 0.8])

    with left:
        st.markdown("### About AI Interviewer")
        items = [
            ("🚀", "Resume-based smart question generation"),
            ("🎯", "Real-time AI evaluation & scoring"),
            ("💡", "Detailed feedback & improvement tips"),
            ("⏱️", "30-minute timed interview session"),
            ("📥", "Downloadable PDF performance report"),
            ("✅", "Instant selection / rejection decision"),
        ]
        for icon, text in items:
            st.markdown(f'<div class="about-item"><span>{icon}</span><span>{text}</span></div>', unsafe_allow_html=True)

    with right:
        st.markdown('<div class="login-card">', unsafe_allow_html=True)
        st.markdown('<div class="login-title">Welcome back 👋</div>', unsafe_allow_html=True)
        st.markdown('<div class="login-sub">Sign in to start your interview</div>', unsafe_allow_html=True)

        # Use st.form for login — inputs render properly in Streamlit
        with st.form("login_form", clear_on_submit=False):
            email = st.text_input("Email Address", placeholder="you@example.com")
            password = st.text_input("Password", type="password", placeholder="Enter your password")
            submit = st.form_submit_button("Sign In →", use_container_width=True)

            if submit:
                if password == APP_PASSWORD:
                    st.session_state.logged = True
                    st.rerun()
                elif not email or not password:
                    st.warning("Please fill in all fields.")
                else:
                    st.error("Incorrect password. Please try again.")

        st.markdown('</div>', unsafe_allow_html=True)

# ================================================================
# MAIN APP
# ================================================================
else:
    with st.sidebar:
        st.markdown("## 🤖 AI Interviewer")
        st.markdown("---")
        st.session_state.mode = st.radio(
            "Navigation",
            ["Interview Assistant", "Chatbot"],
            label_visibility="collapsed"
        )
        st.markdown("---")
        if st.button("🔄 Restart Session", use_container_width=True):
            st.session_state.clear()
            st.rerun()
        st.markdown("<br><br>", unsafe_allow_html=True)
        st.markdown('<div style="color:#475569;font-size:12px;">AI Interviewer v2.0</div>', unsafe_allow_html=True)

    # ---- CHATBOT ----
    if st.session_state.mode == "Chatbot":
        st.markdown('<div class="hero-title" style="font-size:38px;">AI Chat Assistant</div>', unsafe_allow_html=True)
        st.markdown('<div class="hero-sub">Ask anything about interviews, resumes, or career advice.</div>', unsafe_allow_html=True)

        if "chat" not in st.session_state:
            st.session_state.chat = []

        # Chat history
        chat_container = st.container()
        with chat_container:
            if not st.session_state.chat:
                st.markdown('<div style="color:#475569;text-align:center;padding:40px 0;">👋 Start a conversation below</div>', unsafe_allow_html=True)
            for role, msg in st.session_state.chat:
                if role == "user":
                    st.markdown(f'<div class="chat-bubble-user">🧑 {msg}</div>', unsafe_allow_html=True)
                else:
                    st.markdown(f'<div class="chat-bubble-ai">🤖 {msg}</div>', unsafe_allow_html=True)

        # Input
        def handle_chat():
            u_msg = st.session_state.chat_input_key
            if u_msg and u_msg.strip():
                reply = chatbot_res(u_msg)
                st.session_state.chat.append(("user", u_msg))
                st.session_state.chat.append(("ai", reply))
                st.session_state.chat_input_key = ""

        st.text_input(
            "Your message",
            key="chat_input_key",
            placeholder="Type your question and press Enter...",
            on_change=handle_chat,
            label_visibility="collapsed"
        )

    # ---- INTERVIEW ----
    else:
        if "step" not in st.session_state:
            st.session_state.step = "upload"
            st.session_state.q = []
            st.session_state.a = []
            st.session_state.e = []
            st.session_state.i = 0

        # ---- UPLOAD ----
        if st.session_state.step == "upload":
            st.markdown('<div class="hero-title" style="font-size:38px;">AI Mock Interview</div>', unsafe_allow_html=True)
            st.markdown('<div class="hero-sub">Upload your resume to generate personalized interview questions.</div>', unsafe_allow_html=True)

            col1, col2 = st.columns([1.4, 1])
            with col1:
                file = st.file_uploader(
                    "Upload Your Resume",
                    type=["pdf", "txt", "docx"],
                    help="Supported formats: PDF, TXT, DOCX"
                )
                if file:
                    st.success(f"✅ Uploaded: **{file.name}**")
                    if st.button("Start Interview →", use_container_width=True):
                        text = extract_text(file)
                        if not text.strip():
                            st.error("Could not extract text from file. Try a different format.")
                        else:
                            st.session_state.q = generate_questions(text)
                            st.session_state.step = "interview"
                            st.session_state.start = time.time()
                            st.rerun()

            with col2:
                st.markdown("""
                <div style="background:rgba(255,255,255,0.04);border:1px solid rgba(255,255,255,0.09);border-radius:14px;padding:20px;">
                    <div style="font-size:14px;font-weight:600;color:#94a3b8;text-transform:uppercase;letter-spacing:0.08em;margin-bottom:14px;">How it works</div>
                    <div style="color:#cbd5e1;font-size:14px;line-height:2;">
                        <div>1️⃣ &nbsp;Upload your resume</div>
                        <div>2️⃣ &nbsp;AI generates 7 questions</div>
                        <div>3️⃣ &nbsp;Answer within 30 minutes</div>
                        <div>4️⃣ &nbsp;Get AI feedback & score</div>
                        <div>5️⃣ &nbsp;Download PDF report</div>
                    </div>
                </div>
                """, unsafe_allow_html=True)

        # ---- INTERVIEW ----
        elif st.session_state.step == "interview":
            elapsed = int(time.time() - st.session_state.start)
            remaining = max(0, 1800 - elapsed)
            mins, secs = remaining // 60, remaining % 60

            if remaining == 0:
                st.session_state.step = "report"
                st.rerun()

            # Header row
            col_timer, col_progress = st.columns([1, 3])
            with col_timer:
                color = "#ef4444" if remaining < 300 else "#facc15"
                st.markdown(f'<div class="timer-badge" style="color:{color};border-color:{color}30;background:rgba({("239,68,68" if remaining < 300 else "250,204,21")},0.1);">⏱ {mins:02d}:{secs:02d}</div>', unsafe_allow_html=True)
            with col_progress:
                st.markdown("<br>", unsafe_allow_html=True)
                progress_val = (st.session_state.i) / 7
                st.progress(progress_val)
                st.caption(f"Question {st.session_state.i + 1} of 7")

            st.markdown("<br>", unsafe_allow_html=True)

            # Question card
            q = st.session_state.q[st.session_state.i]
            st.markdown(f"""
            <div class="question-card">
                <div class="question-num">Question {st.session_state.i + 1}</div>
                <div class="question-text">{q}</div>
            </div>""", unsafe_allow_html=True)

            ans = st.text_area(
                "Your Answer",
                key=f"ans_{st.session_state.i}",
                height=160,
                placeholder="Type your answer here... Be specific and structured in your response."
            )

            col_btn, col_info = st.columns([1, 2])
            with col_btn:
                if st.button("Submit Answer →", use_container_width=True):
                    if not ans.strip():
                        st.warning("⚠️ Please write your answer before submitting.")
                    else:
                        st.session_state.e.append(evaluate(q, ans))
                        st.session_state.a.append(ans)
                        if st.session_state.i < 6:
                            st.session_state.i += 1
                        else:
                            st.session_state.step = "report"
                        st.rerun()
            with col_info:
                st.markdown(f'<div style="color:#475569;font-size:13px;padding-top:12px;">💡 {7 - st.session_state.i - 1} questions remaining</div>', unsafe_allow_html=True)

        # ---- REPORT ----
        elif st.session_state.step == "report":
            st.markdown('<div class="hero-title" style="font-size:38px;">Interview Report</div>', unsafe_allow_html=True)
            st.markdown('<div class="hero-sub">Here\'s how you performed across all 7 questions.</div>', unsafe_allow_html=True)

            # Score calculation
            total_score = 0
            scores = []
            for e in st.session_state.e:
                m = re.search(r"Score:\s*(\d)", e)
                val = int(m.group(1)) if m else 0
                scores.append(val)
                total_score += val

            max_score = len(st.session_state.e) * 5
            pct = int((total_score / max_score) * 100) if max_score > 0 else 0
            passed = sum(1 for s in scores if s >= 4)

            # Score card
            col_score, col_result = st.columns(2)
            with col_score:
                st.markdown(f"""
                <div style="background:rgba(250,204,21,0.08);border:1px solid rgba(250,204,21,0.25);border-radius:16px;padding:24px;text-align:center;">
                    <div style="color:#94a3b8;font-size:13px;font-weight:600;text-transform:uppercase;letter-spacing:0.08em;">Total Score</div>
                    <div class="score-display">{total_score}<span style="font-size:22px;color:#64748b;">/{max_score}</span></div>
                    <div style="color:#94a3b8;font-size:14px;">{pct}% — {passed} of {len(scores)} questions scored 4+</div>
                </div>
                """, unsafe_allow_html=True)
            with col_result:
                if passed >= 5:
                    st.markdown("""
                    <div style="background:rgba(34,197,94,0.1);border:1px solid rgba(34,197,94,0.3);border-radius:16px;padding:24px;text-align:center;">
                        <div style="font-size:42px;">🎉</div>
                        <div style="font-size:22px;font-weight:700;color:#86efac;">Selected</div>
                        <div style="color:#94a3b8;font-size:14px;">Strong performance! You qualify.</div>
                    </div>
                    """, unsafe_allow_html=True)
                else:
                    st.markdown("""
                    <div style="background:rgba(239,68,68,0.1);border:1px solid rgba(239,68,68,0.3);border-radius:16px;padding:24px;text-align:center;">
                        <div style="font-size:42px;">📚</div>
                        <div style="font-size:22px;font-weight:700;color:#fca5a5;">Not Selected</div>
                        <div style="color:#94a3b8;font-size:14px;">Keep practicing and try again!</div>
                    </div>
                    """, unsafe_allow_html=True)

            st.markdown("<br>", unsafe_allow_html=True)

            # PDF Download
            pdf_data = create_pdf(st.session_state.q, st.session_state.e)
            st.download_button(
                "📥 Download Full Report (PDF)",
                pdf_data,
                "AI_Interview_Report.pdf",
                "application/pdf",
                use_container_width=False
            )

            st.markdown("<hr class='section-divider'>", unsafe_allow_html=True)
            st.markdown("### Question-by-Question Breakdown")

            for i, (e, score) in enumerate(zip(st.session_state.e, scores)):
                color = "#22c55e" if score >= 4 else "#ef4444" if score <= 2 else "#f59e0b"
                with st.expander(f"Q{i+1}: {st.session_state.q[i][:70]}...  |  Score: {score}/5", expanded=(i == 0)):
                    col_q, col_s = st.columns([3, 1])
                    with col_q:
                        st.markdown(f"**Question:** {st.session_state.q[i]}")
                        st.markdown(f"**Your Answer:** {st.session_state.a[i]}")
                        st.markdown("---")
                        st.markdown(e)
                    with col_s:
                        st.markdown(f'<div style="font-size:36px;font-weight:700;color:{color};text-align:center;">{score}/5</div>', unsafe_allow_html=True)